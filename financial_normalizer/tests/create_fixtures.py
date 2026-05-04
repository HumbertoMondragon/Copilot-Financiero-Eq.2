"""
Run this script once to generate synthetic test fixtures.

  python financial_normalizer/tests/create_fixtures.py
  python -m financial_normalizer.tests.create_fixtures

Excel structure matches the real-world format:
  - Section header row: "VENTAS" / "Egresos" — label only, no value
  - Group header row:   "A" — single letter with group total in month col
  - Subline rows:       "Alimentos", "Bebidas", ...
  - KPI summary rows at bottom using the exact labels from the real file,
    with percentage values written as "XX.XX%" strings (e.g. "35.84%")
"""
import sys
from pathlib import Path

if __package__ is None:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from financial_normalizer.tests.fixtures import (
    VENTAS, EGRESOS,
    TOTAL_VENTAS, TOTAL_COSTO, MARGEN_COSTO,
    UTILIDAD_BRUTA, MARGEN_BRUTO,
    NOMINA, GASTOS_COMERCIALES, GASTOS_OPERATIVOS, GASTOS_ADMIN, VIATICOS,
    TOTAL_GASTOS_OP, MARGEN_OPERACION,
    EBITDA, MARGEN_EBITDA,
    GASTOS_FINANCIEROS, COSTO_INTEGRAL_FINANCIAMIENTO,
    IMPUESTOS, UTILIDAD_NETA, MARGEN_NETO,
    PROSE_PDF_TEXT,
)

import openpyxl
from openpyxl.styles import Font

FIXTURES_DIR = Path(__file__).parent / "data"
SUBLINES = ["Alimentos", "Bebidas", "Cafe_Inf", "Destilados", "Vinos"]


def _pct(ratio: float) -> str:
    """Convert a ratio (0.3584) to a percentage string ('35.84%')."""
    return f"{ratio * 100:.2f}%"


def create_excel_fixture():
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estado de Resultados"

    bold = Font(bold=True)
    row = [1]

    def wr(label, value=None, is_bold=False):
        ws.cell(row[0], 1, label)
        if value is not None:
            ws.cell(row[0], 2, value)
        if is_bold:
            ws.cell(row[0], 1).font = bold
        row[0] += 1

    def blank():
        row[0] += 1

    # Header
    ws.cell(1, 1, "Concepto").font = bold
    ws.cell(1, 2, "ENERO 2026").font = bold
    row[0] = 2

    # VENTAS section
    wr("VENTAS", is_bold=True)
    for grp in "ABCDE":
        g = VENTAS[grp]
        wr(grp, g["total"], is_bold=True)
        for sub in SUBLINES:
            wr(sub, g[sub])
        blank()

    # EGRESOS section
    wr("Egresos", is_bold=True)
    for grp in "ABCDE":
        g = EGRESOS[grp]
        wr(grp, g["total"], is_bold=True)
        for sub in SUBLINES:
            wr(sub, g[sub])
        blank()

    # KPI summary rows — exact labels as they appear in the real file
    # Percentage values are written as "XX.XX%" strings so the parser can detect them
    wr("Total COSTO DE LO VENDIDO", TOTAL_COSTO, is_bold=True)
    wr("% venta", _pct(MARGEN_COSTO))                    # occurrence 0 → margen_costo
    wr("UTILIDAD BRUTA", UTILIDAD_BRUTA, is_bold=True)
    wr("Margen", _pct(MARGEN_BRUTO))                     # occurrence 0 → margen_bruto
    blank()
    wr("NÓMINA", NOMINA)
    wr("GASTOS COMERCIALES", GASTOS_COMERCIALES)
    wr("GASTOS OPERATIVOS", GASTOS_OPERATIVOS)
    wr("GASTOS ADMINISTRATIVOS", GASTOS_ADMIN)
    wr("VIÁTICOS", VIATICOS)
    wr("Total GASTOS DE OPERACIÓN", TOTAL_GASTOS_OP, is_bold=True)
    wr("% venta", _pct(MARGEN_OPERACION))                # occurrence 1 → margen_operacion
    blank()
    wr("EBITDA", EBITDA, is_bold=True)
    wr("Margen", _pct(MARGEN_EBITDA))                    # occurrence 1 → margen_ebitda
    blank()
    wr("GASTOS FINANCIEROS", GASTOS_FINANCIEROS)
    wr("COSTO INTEGRAL DE FINANCIAMIENTO", COSTO_INTEGRAL_FINANCIAMIENTO)
    wr("IMPUESTOS", IMPUESTOS)
    wr("Utilidad (ó Pérdida)", UTILIDAD_NETA, is_bold=True)
    wr("Margen", _pct(MARGEN_NETO))                      # occurrence 2 → margen_neto

    out = FIXTURES_DIR / "sample_enero_2026.xlsx"
    wb.save(str(out))
    print(f"Excel fixture written: {out}")


def create_text_fixture():
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    out = FIXTURES_DIR / "sample_enero_2026.txt"
    out.write_text(PROSE_PDF_TEXT, encoding="utf-8")
    print(f"Text fixture written: {out}")


if __name__ == "__main__":
    create_excel_fixture()
    create_text_fixture()
